from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parent
OUTPUT = ROOT / "Task-List-Deployment.pptx"
DECK = Presentation()
DECK.slide_width = Inches(13.333)
DECK.slide_height = Inches(7.5)
DECK.core_properties.title = "From Code to Cloud: Task List Deployment"
DECK.core_properties.subject = "A beginner-friendly GitHub Actions team presentation"
DECK.core_properties.author = "Task List Team"
COLORS = {
    "paper": "F5F7FA", "white": "FFFFFF", "ink": "172124",
    "muted": "526168", "blue": "1569D4", "green": "16745A",
    "line": "D9E1E8", "lightblue": "E8F0FC", "lightgreen": "E5F2EC",
    "code": "202B30", "gold": "AA6B0A", "lightgold": "FBF0D9",
}


def color(name):
    return RGBColor.from_string(COLORS.get(name, name))


def box(slide, x, y, width, height, fill, line=None, kind=MSO_SHAPE.RECTANGLE):
    shape = slide.shapes.add_shape(
        kind, Inches(x), Inches(y), Inches(width), Inches(height)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = color(fill)
    if line:
        shape.line.color.rgb = color(line)
    else:
        shape.line.fill.background()
    return shape


def text(slide, content, x, y, width, height, size=22, fill="ink", bold=False,
         font="Avenir Next"):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(width), Inches(height))
    frame = shape.text_frame
    frame.word_wrap = True
    frame.margin_left = frame.margin_right = Inches(0)
    frame.margin_top = frame.margin_bottom = Inches(0)
    for index, line in enumerate(content.split("\n")):
        paragraph = frame.paragraphs[0] if index == 0 else frame.add_paragraph()
        paragraph.text = line
        paragraph.font.name = font
        paragraph.font.size = Pt(size)
        paragraph.font.bold = bold
        paragraph.font.color.rgb = color(fill)
        paragraph.space_after = Pt(10)
    return shape


def base(section, title, subtitle, number, notes):
    slide = DECK.slides.add_slide(DECK.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = color("paper")
    box(slide, 0, 0, 0.16, 7.5, "green")
    text(slide, section.upper(), 0.65, 0.45, 11.8, 0.3, 12, "green", True)
    text(slide, title, 0.65, 0.98, 12, 0.75, 34, bold=True)
    text(slide, subtitle, 0.65, 1.87, 11.9, 0.65, 18, "muted")
    box(slide, 0.65, 6.94, 12, 0.015, "line")
    text(slide, "TASK LIST  /  TEAM WALKTHROUGH", 0.65, 7.1, 8, 0.25, 10, "muted")
    text(slide, f"{number:02d} / 08", 11.65, 7.08, 1, 0.28, 11, "muted")
    slide.notes_slide.notes_text_frame.text = notes
    return slide


def label(slide, tag, title, description, x, y, width=2.75, accent="blue"):
    box(slide, x, y, width, 0.05, accent)
    text(slide, tag, x, y + 0.2, width, 0.4, 13, accent, True)
    text(slide, title, x, y + 0.76, width, 0.6, 24, bold=True)
    text(slide, description, x, y + 1.5, width, 1.25, 18, "muted")


def code(slide, content, x, y, width, height, size=17):
    box(slide, x, y, width, height, "code")
    text(slide, content, x + 0.25, y + 0.25, width - 0.5, height - 0.5,
         size, "white", font="Menlo")


def band(slide, title, detail, y=5.82, fill="lightgreen", accent="green"):
    box(slide, 0.65, y, 12, 0.82, fill)
    text(slide, title, 0.9, y + 0.15, 2.5, 0.45, 16, accent, True)
    text(slide, detail, 3.35, y + 0.15, 9, 0.5, 17)


slide = base(
    "From code to cloud", "One push. A repeatable deployment.",
    "How Task List reaches Azure App Service with GitHub Actions.", 1,
    "Introduce the goal: a small code change should reach our website without manual file copying. "
    "GitHub Actions builds and deploys; Azure runs the app. This deck explains our actual dev workflow, "
    "not a Docker pipeline. Allow about five to seven minutes for the walkthrough."
)
for index, (tag, title, description, accent) in enumerate([
    ("01 / SOURCE", "Change", "Update the app\nand push to dev.", "ink"),
    ("02 / GITHUB", "Build", "Prepare the files\nAzure can run.", "blue"),
    ("03 / AZURE", "Deploy", "Upload files to\nApp Service.", "green"),
    ("04 / WEBSITE", "Use", "Open the updated\nTask List app.", "ink"),
]):
    label(slide, tag, title, description, 0.65 + index * 3.1, 3.0, 2.55, accent)
    if index < 3:
        box(slide, 3.36 + index * 3.1, 3.84, 0.23, 0.23, "muted", kind=MSO_SHAPE.CHEVRON)
band(slide, "OUR APPROACH", ".NET 10 code deployment. No Docker image or ACR required.")

slide = base(
    "01 / Start", "What starts the workflow?", "Automatic when we push. Manual when we need it.", 2,
    "Show the push trigger for the dev branch. Other branches do not trigger this workflow on push. "
    "workflow_dispatch enables the Run workflow button in GitHub Actions. The job uses the GitHub "
    "environment named dev, which holds the publish-profile secret and can enforce approvals. "
    "The concurrency group happens to be called azure-production; that is a lock name, not the target environment."
)
label(slide, "AUTOMATIC", "Push to dev", "A commit pushed to dev\nstarts Build and deploy.", 0.65, 2.85, 3.3)
label(slide, "MANUAL", "Run workflow", "Open GitHub Actions\nand start it yourself.", 4.45, 2.85, 3.3, "green")
code(slide, "on:\n  push:\n    branches: [dev]\n  workflow_dispatch:", 8.45, 2.85, 4.2, 2.65)
band(slide, "ENVIRONMENT", "The job uses dev for its GitHub deployment secret.")

slide = base(
    "02 / Prepare", "GitHub gives us a build machine.",
    "A runner is simply a temporary computer that executes our steps.", 3,
    "Explain runs-on ubuntu-latest: GitHub supplies a fresh Linux machine for the job. "
    "Checkout downloads the repository into that machine. setup-dotnet installs a .NET 10 SDK. "
    "The SDK is needed to compile here; Azure App Service only needs the .NET runtime to run the published app. "
    "The runner is separate from both our laptop and the Azure server."
)
box(slide, 0.65, 2.85, 4.0, 2.55, "lightblue")
text(slide, "GITHUB-HOSTED RUNNER", 0.95, 3.1, 3.4, 0.4, 13, "blue", True)
text(slide, "Temporary\nUbuntu machine", 0.95, 3.65, 3.4, 1.25, 28, bold=True)
label(slide, "STEP 1", "Download code", "actions/checkout@v4\nfetches our repository.", 5.2, 2.85, 3.1)
label(slide, "STEP 2", "Install .NET", "actions/setup-dotnet@v4\ninstalls the .NET 10 SDK.", 9.0, 2.85, 3.3, "green")
band(slide, "GOOD TO KNOW", "The runner builds the app; it does not host our website.")

slide = base(
    "03 / Build", "Publish turns code into runnable files.",
    "We send Azure the finished application, not the source project.", 4,
    "Walk through dotnet publish: it restores packages, builds in Release configuration and collects "
    "deployable files. The complete workflow command includes the output path under runner.temp/task-app. "
    "The package includes TaskApp.dll, runtime configuration, dependency files and wwwroot assets. "
    "set -euo pipefail stops the shell on failures. If publishing fails, deployment does not run. "
    "There is currently no separate automated test step or test project."
)
code(slide, "dotnet publish TaskApp.csproj\n  --configuration Release\n  --output \"${{ runner.temp }}/task-app\"", 0.65, 2.85, 7.0, 2.35, 17)
text(slide, "WHAT THIS DOES", 0.65, 5.36, 2.8, 0.35, 13, "blue", True)
text(slide, "Restore dependencies  /  Compile  /  Collect files", 0.65, 5.84, 7.0, 0.7, 19)
box(slide, 8.25, 2.85, 4.4, 3.5, "white", "line")
text(slide, "task-app/", 8.55, 3.1, 3.8, 0.45, 22, "green", True, "Menlo")
text(slide, "TaskApp.dll\nDependencies\nRuntime configuration\nwwwroot/ assets", 8.55, 3.85, 3.8, 2.05, 19)

slide = base(
    "04 / Access", "The publish profile is our deployment key.",
    "It lets this workflow upload files to one App Service.", 5,
    "Explain that a publish profile is downloaded XML containing deployment credentials; it is sensitive. "
    "Enable SCM Basic Auth Publishing Credentials in the web app's Configuration, General settings. "
    "FTP basic authentication is not required. Download the profile and store the entire XML in the GitHub "
    "dev environment secret named AZURE_WEBAPP_PUBLISH_PROFILE. Do not show a real profile during the demo. "
    "This workflow does not use azure/login or OIDC. If an organizational policy blocks basic auth, "
    "an administrator must approve a supported authentication approach."
)
for index, (tag, title, description, accent) in enumerate([
    ("AZURE PORTAL", "Download", "Enable SCM basic auth.\nDownload the profile.", "blue"),
    ("GITHUB / DEV", "Store securely", "Save the entire XML\nas an environment secret.", "green"),
    ("DEPLOYMENT STEP", "Authenticate", "The action uses the secret\nto upload the app.", "ink"),
]):
    label(slide, tag, title, description, 0.65 + index * 4.1, 2.8, 3.55, accent)
code(slide, "AZURE_WEBAPP_PUBLISH_PROFILE", 0.65, 5.24, 12, 0.9, 20)
text(slide, "Treat it like a password. Never commit it, display it, or share it.",
     0.65, 6.19, 12, 0.4, 17, "gold", True)

slide = base(
    "05 / Deploy", "Upload the app. Azure runs it.",
    "Our target is tododevday-app on Linux Azure App Service.", 6,
    "The azure/webapps-deploy action uploads the task-app output folder using the publish profile. "
    "App Service runs the configured startup command dotnet TaskApp.dll on its built-in .NET 10 runtime. "
    "The app becomes ready after startup, not instantly when the upload starts. SQLite is stored in /home/data, "
    "outside the deployed app files, so redeployment does not intentionally replace the database. "
    "Keep a single App Service instance with this SQLite design and back up the data."
)
label(slide, "GITHUB ACTION", "Upload package", "azure/webapps-deploy@v3\nsends the published files.", 0.65, 2.85, 3.6)
label(slide, "AZURE RUNTIME", "Start the app", "dotnet TaskApp.dll\nlaunches the website.", 4.9, 2.85, 3.3, "green")
label(slide, "USER EXPERIENCE", "Open Task List", "The updated application\nis available in the browser.", 9.0, 2.85, 3.5, "ink")
band(slide, "TASK DATA", "SQLite stays separately in /home/data on persistent storage.")

slide = base(
    "06 / Confidence", "Know when deployment worked.",
    "Watch the workflow, then check the running application.", 7,
    "Show the Actions page and expand a step's logs. Steps execute in sequence and a failed step normally "
    "prevents subsequent steps from running. A successful upload should be followed by a live app check. "
    "Open /health and look for Healthy, then verify the visible change on the home page. The workflow does "
    "not automatically call /health; this is a manual verification step. Its concurrency group serializes "
    "deployments and cancel-in-progress false allows a running deployment to finish."
)
label(slide, "1 / ACTIONS", "Check the run", "Open Build and deploy.\nInspect each step's result.", 0.65, 2.85, 3.5)
label(slide, "2 / HEALTH", "Check startup", "Open /health.\nExpected response: Healthy.", 4.8, 2.85, 3.5, "green")
label(slide, "3 / WEBSITE", "Check the change", "Refresh the app.\nConfirm the updated heading.", 8.95, 2.85, 3.5, "ink")
band(slide, "GUARDRAILS", "Failures stop the job. Only one deployment runs at a time.")

slide = base(
    "Team demo", "Make one change. Follow it to Azure.",
    "A small visible change makes the whole journey easy to follow.", 8,
    "Use a harmless heading change on the dev branch. Commit and push, open the Actions run, show checkout, "
    "SDK setup, publish and deploy, then refresh the website after the deployment succeeds. Avoid changing "
    "or deleting task data during the demo. Close with the main idea: a push starts an automated process "
    "that builds the code and uploads the result to Azure using a securely stored publish profile. "
    "Allow time for questions and for Azure startup."
)
for index, (title, detail) in enumerate([
    ("Edit a heading", "Make a small change in the app."),
    ("Commit and push", "Push the update to the dev branch."),
    ("Watch GitHub Actions", "Follow the build and deployment steps."),
    ("Refresh the website", "Confirm that the new heading is visible."),
]):
    position = 2.75 + index * 0.75
    box(slide, 0.65, position, 0.43, 0.43, "blue" if index < 2 else "green")
    text(slide, str(index + 1), 0.79, position + 0.025, 0.25, 0.35, 16, "white", True)
    text(slide, title, 1.35, position, 3.4, 0.48, 20, bold=True)
    text(slide, detail, 5.0, position + 0.015, 7.5, 0.48, 18, "muted")
text(slide, "tododevday-app.azurewebsites.net", 0.65, 6.08, 11.9, 0.5, 23, "blue", True)


def validate_bounds():
    for index, slide in enumerate(DECK.slides, 1):
        for shape in slide.shapes:
            if (shape.left < 0 or shape.top < 0
                    or shape.left + shape.width > DECK.slide_width
                    or shape.top + shape.height > DECK.slide_height):
                raise ValueError(f"Slide {index}: shape outside slide bounds")
        if not slide.notes_slide.notes_text_frame.text.strip():
            raise ValueError(f"Slide {index}: presenter notes missing")


if __name__ == "__main__":
    validate_bounds()
    DECK.save(OUTPUT)
    print(f"Created {OUTPUT} ({len(DECK.slides)} slides, with presenter notes)")